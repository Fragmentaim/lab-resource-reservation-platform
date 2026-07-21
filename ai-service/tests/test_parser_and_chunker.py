from docx import Document
from openpyxl import Workbook

from app.core.chunker import chunk_document
from app.core.document_ir import DocumentBlock, ParseQualityReport, ParsedDocument, SourceLocation
from app.core.parser import (
    PdfParserProvider,
    SpreadsheetParserProvider,
    WordParserProvider,
    _pdf_page_decision,
)


def test_pdf_page_router_distinguishes_native_hybrid_ocr_and_review():
    assert _pdf_page_decision("原生文本" * 50, 0.0, 0)[0] == "NATIVE"
    assert _pdf_page_decision("原生文本" * 50, 0.25, 1)[0] == "HYBRID"
    assert _pdf_page_decision("", 0.50, 1)[0] == "OCR"
    assert _pdf_page_decision("过短文本", 0.0, 0)[0] == "REVIEW"


def test_word_and_spreadsheet_parsers_keep_structured_metadata(tmp_path):
    word_path = tmp_path / "policy.docx"
    document = Document()
    document.add_heading("预约规则", level=1)
    document.add_paragraph("预约需要遵守实验室制度。")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "字段"
    table.cell(0, 1).text = "说明"
    table.cell(1, 0).text = "时段"
    table.cell(1, 1).text = "不可重叠"
    document.save(word_path)

    parsed_word = WordParserProvider().parse(str(word_path))
    text_block = next(block for block in parsed_word.blocks if block.block_type == "TEXT")
    assert text_block.title_path == ["预约规则"]
    assert any(block.block_type == "TABLE" for block in parsed_word.blocks)

    sheet_path = tmp_path / "resources.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "资源"
    sheet.append(["资源", "状态"])
    sheet.append(["靶车", "可预约"])
    workbook.save(sheet_path)

    parsed_sheet = SpreadsheetParserProvider().parse(str(sheet_path))
    assert parsed_sheet.blocks[0].location.unit_label == "sheet:资源"
    assert parsed_sheet.blocks[0].title_path == ["资源"]
    assert "| 资源 | 状态 |" in parsed_sheet.blocks[0].content


def test_chunker_carries_heading_path_source_location_and_hash():
    parsed = ParsedDocument(
        blocks=[DocumentBlock(
            "TEXT",
            "# 实验室制度\n预约前需要确认资源状态。",
            SourceLocation(2, "page:2"),
        )],
        quality=ParseQualityReport(
            provider="test-provider",
            provider_version="1",
            parse_mode="TEXT",
            unit_count=1,
            non_empty_unit_count=1,
            character_count=20,
        ),
    )

    chunks = chunk_document(parsed)

    assert len(chunks) == 1
    assert chunks[0]["title_path"] == ["实验室制度"]
    assert chunks[0]["source_location"]["unit_label"] == "page:2"
    assert chunks[0]["parser_provider"] == "test-provider"
    assert chunks[0]["content_hash"]
    assert chunks[0]["token_count"] > 0


def test_chunker_keeps_structured_blocks_atomic_and_marks_their_strategy():
    parsed = ParsedDocument(
        blocks=[
            DocumentBlock(
                "TABLE",
                "| 命令码 | 说明 |\n| --- | --- |\n| 0x0101 | 图形数据 |",
                SourceLocation(7, "page:7"),
                ["串口协议"],
            ),
            DocumentBlock(
                "CODE",
                "typedef struct {\n  uint16_t cmd_id;\n} frame_t;",
                SourceLocation(7, "page:7"),
                ["串口协议"],
            ),
        ],
        quality=ParseQualityReport(
            provider="test-provider",
            provider_version="1",
            parse_mode="STRUCTURED",
            unit_count=1,
            non_empty_unit_count=1,
            character_count=100,
        ),
    )

    chunks = chunk_document(parsed)

    assert len(chunks) == 2
    assert [chunk["block_type"] for chunk in chunks] == ["TABLE", "CODE"]
    assert all(chunk["chunk_strategy"] == "ATOMIC" for chunk in chunks)
    assert "0x0101" in chunks[0]["content"]
    assert "frame_t" in chunks[1]["content"]


def test_chunker_detects_fenced_code_without_document_specific_rules():
    parsed = ParsedDocument(
        blocks=[DocumentBlock(
            "TEXT",
            "# 示例\n说明文字。\n```c\ntypedef struct {\n  int id;\n} item_t;\n```\n结束说明。",
            SourceLocation(3, "page:3"),
        )],
        quality=ParseQualityReport(
            provider="test-provider",
            provider_version="1",
            parse_mode="TEXT",
            unit_count=1,
            non_empty_unit_count=1,
            character_count=50,
        ),
    )

    chunks = chunk_document(parsed)

    code_chunks = [chunk for chunk in chunks if chunk["block_type"] == "CODE"]
    assert len(code_chunks) == 1
    assert code_chunks[0]["chunk_strategy"] == "ATOMIC"
    assert "item_t" in code_chunks[0]["content"]
