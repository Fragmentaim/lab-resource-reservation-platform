"""Document parser providers that emit a common Document IR.

The first provider set intentionally uses lightweight, local libraries.  OCR is
optional: a scanned PDF or image is reported as degraded instead of silently
being indexed as an empty document.
"""

from __future__ import annotations

import re
import shutil
from abc import ABC, abstractmethod
from io import BytesIO
from pathlib import Path

import fitz  # pymupdf
import numpy as np
from docx import Document as DocxDocument

from app.config import settings
from app.core.document_ir import (
    DocumentBlock,
    ParseQualityReport,
    ParsedDocument,
    SourceLocation,
)
from app.core.ocr_engine import PaddleOcrEngine


class DocumentParserProvider(ABC):
    name: str
    version: str = "1"
    file_types: set[str]

    def supports(self, file_type: str) -> bool:
        return file_type.upper() in self.file_types

    @abstractmethod
    def parse(self, file_path: str) -> ParsedDocument:
        raise NotImplementedError


class PdfParserProvider(DocumentParserProvider):
    name = "pymupdf-native"
    version = fitz.VersionBind
    file_types = {"PDF"}

    def parse(self, file_path: str) -> ParsedDocument:
        blocks: list[DocumentBlock] = []
        warnings: list[str] = []
        scanned_units = 0
        empty_units = 0
        image_count = 0
        ocr_units = 0
        ocr_characters = 0
        ocr_scores: list[float] = []
        decisions: list[dict] = []
        doc = fitz.open(file_path)
        try:
            for index, page in enumerate(doc, start=1):
                text = page.get_text("text").strip()
                image_regions = _pdf_image_regions(page)
                page_images = len(image_regions)
                image_coverage = _image_coverage(page, image_regions)
                image_count += page_images
                decision, reasons = _pdf_page_decision(text, image_coverage, page_images)
                decisions.append({
                    "unit_no": index,
                    "unit_label": f"page:{index}",
                    "route": decision,
                    "reasons": reasons,
                    "native_character_count": len(text),
                    "image_coverage": round(image_coverage, 4),
                })
                if decision in {"NATIVE", "HYBRID", "REVIEW"} and text:
                    blocks.append(DocumentBlock(
                        block_type="TEXT",
                        content=text,
                        location=SourceLocation(index, f"page:{index}"),
                        metadata={"parse_route": decision},
                    ))
                if decision == "HYBRID":
                    try:
                        image_blocks, image_chars, image_scores = self._ocr_image_regions(
                            page, index, image_regions
                        )
                        blocks.extend(image_blocks)
                        ocr_units += len(image_blocks)
                        ocr_characters += image_chars
                        ocr_scores.extend(image_scores)
                    except RuntimeError as exc:
                        warnings.append(f"PAGE_{index}_HYBRID_OCR_FAILED")
                        decisions[-1]["ocr_error"] = str(exc)
                elif decision == "OCR":
                    scanned_units += 1
                    try:
                        page_block, chars, score = self._ocr_page(page, index)
                    except RuntimeError as exc:
                        page_block, chars, score = None, 0, None
                        decisions[-1]["ocr_error"] = str(exc)
                    if page_block is None:
                        empty_units += 1
                        warnings.append(f"PAGE_{index}_OCR_FAILED_OR_EMPTY")
                    else:
                        blocks.append(page_block)
                        ocr_units += 1
                        ocr_characters += chars
                        if score is not None:
                            ocr_scores.append(score)
                elif decision == "REVIEW":
                    warnings.append(f"PAGE_{index}_REVIEW_RECOMMENDED")
                    if not text:
                        empty_units += 1
            character_count = sum(len(block.content) for block in blocks)
            if ocr_units:
                warnings.append("OCR_APPLIED_TO_SELECTED_CONTENT")
            quality = _quality(
                provider=self,
                parse_mode="NATIVE_TEXT",
                unit_count=len(doc),
                blocks=blocks,
                character_count=character_count,
                image_count=image_count,
                scanned_units=scanned_units,
                ocr_units=ocr_units,
                ocr_characters=ocr_characters,
                ocr_average_confidence=(
                    sum(ocr_scores) / len(ocr_scores) if ocr_scores else None
                ),
                empty_units=empty_units,
                warnings=warnings,
                decisions=decisions,
            )
            return ParsedDocument(blocks=blocks, quality=quality)
        finally:
            doc.close()

    def _ocr_page(self, page, page_no: int) -> tuple[DocumentBlock | None, int, float | None]:
        result = _ocr_pixmap(page.get_pixmap(matrix=fitz.Matrix(2, 2), alpha=False))
        if not result.text:
            return None, 0, result.average_confidence
        return (
            DocumentBlock(
                "TEXT",
                result.text,
                SourceLocation(page_no, f"page:{page_no}"),
                metadata={
                    "parse_route": "OCR",
                    "ocr_provider": PaddleOcrEngine.provider,
                    "ocr_confidence": result.average_confidence,
                },
            ),
            len(result.text),
            result.average_confidence,
        )

    def _ocr_image_regions(self, page, page_no: int, regions: list[fitz.Rect]) -> tuple[list[DocumentBlock], int, list[float]]:
        blocks: list[DocumentBlock] = []
        characters = 0
        scores: list[float] = []
        for image_index, rect in enumerate(regions, start=1):
            if rect.get_area() / page.rect.get_area() < settings.ocr_hybrid_image_coverage:
                continue
            result = _ocr_pixmap(page.get_pixmap(
                matrix=fitz.Matrix(2, 2), clip=rect, alpha=False
            ))
            if not result.text or (result.average_confidence is not None and result.average_confidence < 0.55):
                continue
            blocks.append(DocumentBlock(
                "IMAGE",
                result.text,
                SourceLocation(page_no, f"page:{page_no}/image:{image_index}"),
                metadata={
                    "parse_route": "HYBRID",
                    "ocr_provider": PaddleOcrEngine.provider,
                    "ocr_confidence": result.average_confidence,
                    "bbox": [round(v, 1) for v in rect],
                },
            ))
            characters += len(result.text)
            if result.average_confidence is not None:
                scores.append(result.average_confidence)
        return blocks, characters, scores


class WordParserProvider(DocumentParserProvider):
    name = "python-docx-structured"
    version = "1"
    file_types = {"DOCX"}

    def parse(self, file_path: str) -> ParsedDocument:
        doc = DocxDocument(file_path)
        blocks: list[DocumentBlock] = []
        title_path: list[str] = []
        heading_count = 0

        for paragraph_index, paragraph in enumerate(doc.paragraphs, start=1):
            text = paragraph.text.strip()
            if not text:
                continue
            heading_level = _docx_heading_level(paragraph)
            location = SourceLocation(1, f"paragraph:{paragraph_index}")
            if heading_level:
                heading_count += 1
                title_path = title_path[:heading_level - 1] + [text]
                blocks.append(DocumentBlock("TITLE", text, location, list(title_path)))
            else:
                blocks.append(DocumentBlock("TEXT", text, location, list(title_path)))

        for table_index, table in enumerate(doc.tables, start=1):
            markdown = _table_to_markdown([
                [cell.text.strip() for cell in row.cells]
                for row in table.rows
            ])
            if markdown:
                blocks.append(DocumentBlock(
                    "TABLE",
                    markdown,
                    SourceLocation(1, f"table:{table_index}"),
                    list(title_path),
                    {"table_index": table_index},
                ))

        character_count = sum(len(block.content) for block in blocks)
        quality = _quality(
            provider=self,
            parse_mode="STRUCTURED",
            unit_count=1,
            blocks=blocks,
            character_count=character_count,
            heading_count=heading_count,
            table_count=len(doc.tables),
            empty_units=0 if blocks else 1,
            warnings=[] if blocks else ["DOCX_NO_EXTRACTABLE_CONTENT"],
        )
        return ParsedDocument(blocks=blocks, quality=quality)


class SpreadsheetParserProvider(DocumentParserProvider):
    name = "openpyxl-sheet"
    version = "1"
    file_types = {"XLSX"}
    max_rows_per_sheet = 500
    max_columns_per_sheet = 80

    def parse(self, file_path: str) -> ParsedDocument:
        try:
            from openpyxl import load_workbook
        except ImportError as exc:
            raise RuntimeError("XLSX parsing requires openpyxl. Install ai-service requirements first.") from exc

        workbook = load_workbook(file_path, read_only=True, data_only=True)
        blocks: list[DocumentBlock] = []
        warnings: list[str] = []
        try:
            for sheet_index, sheet in enumerate(workbook.worksheets, start=1):
                max_row = min(sheet.max_row or 0, self.max_rows_per_sheet)
                max_column = min(sheet.max_column or 0, self.max_columns_per_sheet)
                if (sheet.max_row or 0) > max_row or (sheet.max_column or 0) > max_column:
                    warnings.append(
                        f"SHEET_{sheet.title}_TRUNCATED_TO_{max_row}x{max_column}"
                    )
                rows = [
                    ["" if value is None else _normalize_cell(value) for value in row]
                    for row in sheet.iter_rows(
                        min_row=1, max_row=max_row, min_col=1, max_col=max_column, values_only=True
                    )
                ]
                while rows and not any(rows[-1]):
                    rows.pop()
                if not rows:
                    continue
                markdown = _table_to_markdown(rows)
                if markdown:
                    blocks.append(DocumentBlock(
                        "TABLE",
                        markdown,
                        SourceLocation(sheet_index, f"sheet:{sheet.title}"),
                        [sheet.title],
                        {"sheet_name": sheet.title, "row_count": len(rows)},
                    ))
            character_count = sum(len(block.content) for block in blocks)
            quality = _quality(
                provider=self,
                parse_mode="STRUCTURED",
                unit_count=len(workbook.worksheets),
                blocks=blocks,
                character_count=character_count,
                table_count=len(blocks),
                empty_units=max(0, len(workbook.worksheets) - len(blocks)),
                warnings=warnings or ([] if blocks else ["XLSX_NO_EXTRACTABLE_CONTENT"]),
            )
            return ParsedDocument(blocks=blocks, quality=quality)
        finally:
            workbook.close()


class TextParserProvider(DocumentParserProvider):
    name = "utf8-text"
    version = "1"
    file_types = {"MD", "TXT", "MARKDOWN"}

    def parse(self, file_path: str) -> ParsedDocument:
        text = Path(file_path).read_text(encoding="utf-8").strip()
        blocks = [DocumentBlock("TEXT", text, SourceLocation(1, "text:1"))] if text else []
        return ParsedDocument(
            blocks=blocks,
            quality=_quality(
                provider=self,
                parse_mode="TEXT",
                unit_count=1,
                blocks=blocks,
                character_count=len(text),
                empty_units=0 if text else 1,
                warnings=[] if text else ["TEXT_NO_EXTRACTABLE_CONTENT"],
            ),
        )


class ImageOcrParserProvider(DocumentParserProvider):
    """Optional Tesseract adapter; absence is an explicit low-quality outcome."""

    name = "tesseract-optional"
    version = "1"
    file_types = {"PNG", "JPG", "JPEG", "BMP", "TIFF"}

    def parse(self, file_path: str) -> ParsedDocument:
        if not shutil.which("tesseract"):
            return _ocr_unavailable_document(self, "OCR_ENGINE_NOT_CONFIGURED")
        try:
            import pytesseract
            from PIL import Image
        except ImportError:
            return _ocr_unavailable_document(self, "OCR_PYTHON_DEPENDENCY_NOT_INSTALLED")

        text = pytesseract.image_to_string(Image.open(file_path), lang="chi_sim+eng").strip()
        blocks = [DocumentBlock("TEXT", text, SourceLocation(1, "image:1"))] if text else []
        return ParsedDocument(
            blocks=blocks,
            quality=_quality(
                provider=self,
                parse_mode="OCR",
                unit_count=1,
                blocks=blocks,
                character_count=len(text),
                image_count=1,
                scanned_units=1,
                empty_units=0 if text else 1,
                warnings=[] if text else ["OCR_NO_TEXT_RECOGNIZED"],
            ),
        )


_PROVIDERS: list[DocumentParserProvider] = [
    PdfParserProvider(),
    WordParserProvider(),
    SpreadsheetParserProvider(),
    TextParserProvider(),
    ImageOcrParserProvider(),
]


def parse_file(file_path: str, file_type: str) -> ParsedDocument:
    path = Path(file_path)
    if not path.exists():
        raise FileNotFoundError(f"File not found: {file_path}")
    normalized_type = file_type.upper().lstrip(".")
    for provider in _PROVIDERS:
        if provider.supports(normalized_type):
            return provider.parse(str(path))
    raise ValueError(f"Unsupported file type: {normalized_type}")


def _quality(
    provider: DocumentParserProvider,
    parse_mode: str,
    unit_count: int,
    blocks: list[DocumentBlock],
    character_count: int,
    heading_count: int = 0,
    table_count: int = 0,
    image_count: int = 0,
    scanned_units: int = 0,
    ocr_units: int = 0,
    ocr_characters: int = 0,
    ocr_average_confidence: float | None = None,
    empty_units: int = 0,
    warnings: list[str] | None = None,
    decisions: list[dict] | None = None,
) -> ParseQualityReport:
    warning_list = warnings or []
    non_empty_units = len({block.location.unit_no for block in blocks if block.content.strip()})
    score = 1.0
    if not blocks:
        score = 0.0
    elif scanned_units and not ocr_units:
        score = min(score, 0.45)
    elif ocr_average_confidence is not None:
        score = min(score, max(0.55, ocr_average_confidence))
    elif warning_list:
        score = min(score, 0.8)
    return ParseQualityReport(
        provider=provider.name,
        provider_version=provider.version,
        parse_mode=parse_mode,
        unit_count=unit_count,
        non_empty_unit_count=non_empty_units,
        character_count=character_count,
        heading_count=heading_count,
        table_count=table_count,
        image_count=image_count,
        scanned_unit_count=scanned_units,
        ocr_unit_count=ocr_units,
        ocr_character_count=ocr_characters,
        ocr_average_confidence=round(ocr_average_confidence, 4) if ocr_average_confidence is not None else None,
        empty_unit_count=empty_units,
        quality_score=score,
        warnings=warning_list,
        decisions=decisions or [],
    )


def _ocr_unavailable_document(provider: DocumentParserProvider, warning: str) -> ParsedDocument:
    return ParsedDocument(
        blocks=[],
        quality=_quality(
            provider=provider,
            parse_mode="OCR_UNAVAILABLE",
            unit_count=1,
            blocks=[],
            character_count=0,
            image_count=1,
            scanned_units=1,
            empty_units=1,
            warnings=[warning],
        ),
        )


def _pdf_image_regions(page) -> list[fitz.Rect]:
    regions: list[fitz.Rect] = []
    for image in page.get_images(full=True):
        for rect in page.get_image_rects(image[0]):
            if rect.get_area() > 0:
                regions.append(rect)
    return regions


def _image_coverage(page, regions: list[fitz.Rect]) -> float:
    page_area = page.rect.get_area()
    if page_area <= 0:
        return 0.0
    return min(1.0, sum(rect.get_area() for rect in regions) / page_area)


def _pdf_page_decision(text: str, image_coverage: float, image_count: int) -> tuple[str, list[str]]:
    native_characters = len(text)
    if native_characters == 0 and image_count:
        return "OCR", ["NO_NATIVE_TEXT", "IMAGE_CONTENT_PRESENT"]
    if (
        image_count
        and image_coverage >= settings.ocr_scan_image_coverage
        and native_characters < settings.ocr_min_native_chars
    ):
        return "OCR", ["NATIVE_TEXT_INSUFFICIENT", "IMAGE_DOMINANT_PAGE"]
    if (
        native_characters > 0
        and image_count
        and image_coverage >= settings.ocr_hybrid_image_coverage
    ):
        return "HYBRID", ["NATIVE_TEXT_PRESENT", "LARGE_IMAGE_REGION_PRESENT"]
    if native_characters >= settings.ocr_min_native_chars:
        return "NATIVE", ["NATIVE_TEXT_SUFFICIENT"]
    return "REVIEW", ["NATIVE_TEXT_INSUFFICIENT", "NO_RELIABLE_OCR_TARGET"]


def _ocr_pixmap(pixmap):
    try:
        from PIL import Image
        image = Image.open(BytesIO(pixmap.tobytes("png"))).convert("RGB")
        return PaddleOcrEngine.shared().recognize(np.array(image))
    except Exception as exc:
        raise RuntimeError(f"PaddleOCR failed: {exc}") from exc


def _docx_heading_level(paragraph) -> int | None:
    style_name = paragraph.style.name if paragraph.style is not None else ""
    style_name_lower = style_name.lower()
    if style_name_lower.startswith("heading"):
        suffix = style_name_lower.replace("heading", "").strip()
        return min(max(int(suffix), 1), 6) if suffix.isdigit() else 1
    if style_name.startswith("标题"):
        suffix = style_name.replace("标题", "").strip()
        return min(max(int(suffix), 1), 6) if suffix.isdigit() else 1
    return None


def _table_to_markdown(rows: list[list[str]]) -> str:
    normalized = [[str(cell).replace("\n", " ").strip() for cell in row] for row in rows]
    normalized = [row for row in normalized if any(row)]
    if not normalized:
        return ""
    column_count = max(len(row) for row in normalized)
    normalized = [row + [""] * (column_count - len(row)) for row in normalized]
    header = normalized[0]
    divider = ["---"] * column_count
    body = normalized[1:]
    return "\n".join(
        ["| " + " | ".join(row) + " |" for row in [header, divider, *body]]
    )


def _normalize_cell(value: object) -> str:
    return re.sub(r"\s+", " ", str(value)).strip()
